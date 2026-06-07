#!/usr/bin/env python3
"""
ppt_tag_to_notes.py — 对 PPTX 每页进行四维标签分类，写入幻灯片备注

用法:
    python ppt_tag_to_notes.py <input.pptx> [--output <output.pptx>]

输出:
    修改后的 PPTX 文件，每页备注中包含标签分类结果
"""

import argparse
import io
import os
import sys
from pathlib import Path

# Windows 控制台 UTF-8（避免重复包装）
if sys.platform == "win32" and hasattr(sys.stdout, "buffer") and not isinstance(sys.stdout, io.TextIOWrapper):
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")
    sys.stderr = io.TextIOWrapper(sys.stderr.buffer, encoding="utf-8", errors="replace")

from pptx import Presentation

# 复用已有的分类逻辑
SCRIPT_DIR = Path(__file__).resolve().parent
sys.path.insert(0, str(SCRIPT_DIR))
from ppt_split_classify import (
    get_content_shapes, classify_logic, classify_domain,
    classify_count, classify_chart_type, tags_to_chinese,
)

TAG_NOTES_PREFIX = "[EW-TTS Tags]"
TAG_NOTES_MARKER = "━━━ 四维标签分类 (EW-TTS v2.0) ━━━"


def format_tags_for_notes(tags: dict) -> str:
    """将四维标签格式化为备注文本（中文）"""
    cn = tags_to_chinese(tags)
    lines = [
        TAG_NOTES_MARKER,
        f"  逻辑拓扑: {cn['logic']}",
        f"  内容领域: {cn['domain']}",
        f"  数量规模: {cn['count']}",
        f"  图形类型: {cn['chart-type']}",
        TAG_NOTES_MARKER,
    ]
    return "\n".join(lines)


def has_existing_tags(notes_text: str) -> bool:
    """检查备注中是否已有标签"""
    return TAG_NOTES_MARKER in notes_text


def strip_existing_tags(notes_text: str) -> str:
    """移除已有的标签块"""
    lines = notes_text.split("\n")
    result = []
    in_tag_block = False
    for line in lines:
        if line.strip() == TAG_NOTES_MARKER:
            in_tag_block = not in_tag_block
            continue
        if not in_tag_block:
            result.append(line)
    return "\n".join(result).strip()


def classify_and_write(input_path: str, output_path: str = None):
    """主流程：逐页分类 + 写入备注"""
    input_path = str(Path(input_path).resolve())

    if not os.path.exists(input_path):
        print(f"❌ 文件不存在: {input_path}")
        sys.exit(1)

    prs = Presentation(input_path)
    slide_width = prs.slide_width
    slide_height = prs.slide_height
    n_slides = len(prs.slides)

    print(f"📂 输入: {input_path}")
    print(f"📊 共 {n_slides} 页幻灯片")
    print()

    for i, slide in enumerate(prs.slides):
        idx = i + 1
        content_shapes = get_content_shapes(slide, slide_width, slide_height)

        tags = {
            "logic": classify_logic(slide, content_shapes, slide_width, slide_height),
            "domain": classify_domain(slide),
            "count": classify_count(content_shapes),
            "chart-type": classify_chart_type(slide, content_shapes, slide_width, slide_height),
        }

        # 获取已有备注
        if slide.has_notes_slide:
            existing = slide.notes_slide.notes_text_frame.text
        else:
            existing = ""

        # 移除旧标签块
        clean_notes = strip_existing_tags(existing)

        # 拼接：已有备注 + 新标签
        tag_text = format_tags_for_notes(tags)
        if clean_notes:
            new_notes = clean_notes + "\n\n" + tag_text
        else:
            new_notes = tag_text

        # 写入备注
        if not slide.has_notes_slide:
            slide.notes_slide  # 触发创建
        slide.notes_slide.notes_text_frame.text = new_notes

        print(f"  slide {idx:3d}: "
              f"logic={tags['logic']:12s} "
              f"domain={tags['domain']:14s} "
              f"count={tags['count']:10s} "
              f"chart-type={tags['chart-type']}")

    # 保存
    if not output_path:
        p = Path(input_path)
        output_path = str(p.parent / f"{p.stem}_tagged{p.suffix}")

    prs.save(output_path)
    print()
    print(f"✅ 已保存: {output_path}")

    return output_path


def main():
    parser = argparse.ArgumentParser(
        description="PPT 每页四维标签分类 → 写入幻灯片备注",
    )
    parser.add_argument("input", help="输入 .pptx 文件路径")
    parser.add_argument("--output", "-o", default=None,
                        help="输出文件路径 (默认: <输入文件名>_tagged.pptx)")

    args = parser.parse_args()
    classify_and_write(args.input, args.output)


if __name__ == "__main__":
    main()
